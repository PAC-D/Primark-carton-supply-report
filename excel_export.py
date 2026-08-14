import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from filtering import BASE_COLUMNS

HEADER_BG = "D9E2F3"
TOTAL_BG = "E2EFDA"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _whole(value):
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return value


def _column_width(values):
    longest = max((len(str(v)) for v in values), default=0)
    return min(longest + 2, 40)


def render_excel(df, title, out):
    if isinstance(out, os.PathLike):
        out = os.fspath(out)
    n = len(BASE_COLUMNS)
    columns = df.columns.tolist()
    data = [columns]
    if len(df):
        for row in df.values.tolist():
            data.append(list(row[:n]) + [_whole(v) for v in row[n:]])
        sums = df.sum(numeric_only=True)
        total_row = []
        for c in columns:
            if c in BASE_COLUMNS:
                total_row.append("Total" if c == "Packaging Supplier" else "")
            else:
                total_row.append(_whole(sums[c]))
        data.append(total_row)

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)

    for c, value in enumerate(data[0], start=1):
        cell = ws.cell(row=2, column=c, value=value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(data[1:], start=3):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            if c > n:
                cell.alignment = Alignment(horizontal="right")

    if len(df):
        last = ws.max_row
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=last, column=c)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=TOTAL_BG)

    for c in range(1, len(columns) + 1):
        col = [ws.cell(row=r, column=c).value for r in range(2, ws.max_row + 1)] + [columns[c - 1]]
        ws.column_dimensions[get_column_letter(c)].width = _column_width(col)

    ws.freeze_panes = "A3"
    wb.save(out)