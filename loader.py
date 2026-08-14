import datetime
from pathlib import Path

import openpyxl

from models import Record

SUMMARY_SHEETS = {"2024", "Uniglory Sales Analy"}


def _num(value):
    try:
        if value is None:
            return 0
        return float(value)
    except (TypeError, ValueError):
        return 0


def _parse_sheet(ws):
    header = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    month_cols = []
    for i, v in enumerate(header):
        if isinstance(v, (datetime.datetime, datetime.date)):
            month_cols.append((i + 1, v.year, v.month))
    if not month_cols:
        return None
    factory_col = 3
    supplier_col = 2
    for r in (2, 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        factory_col = next(
            (i + 1 for i, v in enumerate(row)
             if isinstance(v, str) and v.strip().lower() == "factory"),
            factory_col,
        )
        supplier_col = next(
            (i + 1 for i, v in enumerate(row)
             if isinstance(v, str) and v.strip().lower() == "suppliers"),
            supplier_col,
        )
    records = []
    for r in range(3, ws.max_row + 1):
        supplier = str(ws.cell(row=r, column=supplier_col).value or "").strip()
        factory = str(ws.cell(row=r, column=factory_col).value or "").strip()
        if not supplier and not factory:
            continue
        if "total" in (supplier + " " + factory).lower():
            continue
        row_values = [c.value for c in next(ws.iter_rows(min_row=r, max_row=r))]
        extent = max((i + 1 for i, v in enumerate(row_values) if v is not None), default=0)
        for c, year, month in month_cols:
            if c > extent:
                continue
            records.append(Record(
                packaging_supplier="",  # filled by caller
                supplier=supplier,
                factory=factory,
                year=year,
                month=month,
                cartons=_num(row_values[c - 1]),
            ))
    return records


def load_records(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = {}
    for ws in wb.worksheets:
        title = ws.title.strip()
        parts = title.split("-")
        if len(parts) < 2 or title in SUMMARY_SHEETS:
            continue
        sheets[title] = ws

    groups = {}
    for title in sheets:
        parts = [p.strip() for p in title.split("-")]
        supplier = parts[0]
        year = parts[-1]
        if not year.isdigit():
            continue
        is_re = supplier.endswith(" Re")
        base = supplier[:-3].strip() if is_re else supplier
        groups.setdefault((base, int(year)), []).append((title, is_re))

    chosen = []
    for (base, year), group in groups.items():
        re_sheets = [t for t, is_re in group if is_re]
        chosen.append(re_sheets[0] if re_sheets else group[0][0])

    records = []
    warnings = []
    for title in chosen:
        ws = sheets[title]
        supplier_name = title.split("-")[0].strip()
        if supplier_name.endswith(" Re"):
            supplier_name = supplier_name[:-3].strip()
        parsed = _parse_sheet(ws)
        if parsed is None:
            warnings.append(title)
            continue
        for rec in parsed:
            records.append(Record(
                packaging_supplier=supplier_name,
                supplier=rec.supplier,
                factory=rec.factory,
                year=rec.year,
                month=rec.month,
                cartons=rec.cartons,
            ))
    wb.close()
    return records, warnings
