import datetime
from pathlib import Path

import openpyxl


def make_workbook(tmp_path, sheets):
    """sheets: {sheet_title: [(supplier, factory, {(year, month): cartons}), ...]}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        ws.cell(row=1, column=1, value=title)
        ws.cell(row=2, column=1, value="Sl")
        ws.cell(row=2, column=2, value="Suppliers")
        ws.cell(row=2, column=3, value="Factory ")
        months = sorted({m for _, _, data in rows for m in data})
        for j, (y, m) in enumerate(months, start=4):
            ws.cell(row=2, column=j, value=datetime.datetime(y, m, 1))
        for i, (sup, fac, data) in enumerate(rows, start=3):
            ws.cell(row=i, column=1, value=i - 2)
            ws.cell(row=i, column=2, value=sup)
            ws.cell(row=i, column=3, value=fac)
            for (y, m), v in data.items():
                ws.cell(row=i, column=4 + months.index((y, m)), value=v)
    path = tmp_path / "test.xlsx"
    wb.save(path)
    return path
