from loader import load_records
from tests.fixtures import make_workbook


def test_load_basic(tmp_path):
    path = make_workbook(tmp_path, {
        "M&U - 2026": [
            ("PADMA TEXTILES LTD", "Aspire Garments Ltd PJT (24040)", {(2026, 1): 2010, (2026, 2): 9284}),
            ("TEX INTERNATIONAL LTD", "Aboni Knitwear Ltd", {(2026, 1): 500}),
        ],
        "Union-2026": [
            ("CENTRO INTERNATIONAL SOURCING LTD", "APS Apparels Limited (26018)", {(2026, 1): 11537}),
        ],
    })
    records, warnings = load_records(path)
    assert warnings == []
    assert len(records) == 4
    m = {(r.packaging_supplier, r.year, r.month, r.factory): r.cartons for r in records}
    assert m[("M&U", 2026, 1, "Aspire Garments Ltd PJT (24040)")] == 2010
    assert m[("M&U", 2026, 2, "Aspire Garments Ltd PJT (24040)")] == 9284
    assert m[("M&U", 2026, 1, "Aboni Knitwear Ltd")] == 500
    assert m[("Union", 2026, 1, "APS Apparels Limited (26018)")] == 11537


def test_re_replaces_original(tmp_path):
    path = make_workbook(tmp_path, {
        "M&U - 2023": [("A", "Old Factory", {(2023, 1): 100})],
        "M&U Re - 2023": [("A", "New Factory", {(2023, 1): 200})],
    })
    records, _ = load_records(path)
    assert len(records) == 1
    assert records[0].factory == "New Factory"
    assert records[0].cartons == 200


def test_total_fashion_pjt_row_included(tmp_path):
    # Approved consequence: subtotal rows like "Total Fashion PJT" carry a
    # numeric column A and are no longer dropped by name checks.
    path = make_workbook(tmp_path, {
        "M&U - 2026": [
            ("A", "Factory One", {(2026, 1): 100}),
            ("A", "Total", {(2026, 1): 9000}),
        ],
    })
    records, _ = load_records(path)
    assert len(records) == 2


def test_blank_column_a_row_dropped(tmp_path):
    import datetime
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("M&U - 2026")
    ws.append(["Sl", "Suppliers", "Factory ", None, None])
    ws.cell(row=2, column=4, value=datetime.datetime(2026, 1, 1))
    ws.cell(row=3, column=1, value=1)
    ws.cell(row=3, column=2, value="A")
    ws.cell(row=3, column=3, value="Factory One")
    ws.cell(row=3, column=4, value=100)
    # blank column A -> dropped
    ws.cell(row=4, column=2, value="B")
    ws.cell(row=4, column=3, value="Factory Two")
    ws.cell(row=4, column=4, value=200)
    path = tmp_path / "test.xlsx"
    wb.save(path)
    records, _ = load_records(path)
    assert len(records) == 1
    assert records[0].factory == "Factory One"


def test_text_column_a_row_dropped(tmp_path):
    path = make_workbook(tmp_path, {
        "M&U - 2026": [
            ("A", "Factory One", {(2026, 1): 100}),
        ],
    })
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["M&U - 2026"]
    ws.cell(row=3, column=1, value="Sl")
    wb.save(path)
    records, _ = load_records(path)
    assert len(records) == 0


def test_sheet_without_months_warns(tmp_path):
    path = make_workbook(tmp_path, {
        "M&U - 2026": [("A", "F", {(2026, 1): 10})],
        "M&U - 2020": [("A", "F", {})],
    })
    records, warnings = load_records(path)
    assert len(records) == 1
    assert len(warnings) == 1
    assert "M&U - 2020" in warnings[0]


def test_summary_sheets_ignored_silently(tmp_path):
    path = make_workbook(tmp_path, {
        "M&U - 2026": [("A", "F", {(2026, 1): 10})],
        "2024": [("A", "F", {})],
        "Uniglory Sales Analy": [("A", "F", {})],
    })
    records, warnings = load_records(path)
    assert len(records) == 1
    assert warnings == []


def test_empty_cells_are_zero(tmp_path):
    path = make_workbook(tmp_path, {
        "M&U - 2026": [("A", "F", {(2026, 1): None, (2026, 2): 50})],
    })
    records, _ = load_records(path)
    assert len(records) == 2
    by_month = {r.month: r.cartons for r in records}
    assert by_month[1] == 0
    assert by_month[2] == 50


def test_epyllion_layout_factory_in_column_d(tmp_path):
    # Epyllion sheets have an extra "Supplier" column: Sl | Suppliers | Supplier | Factory | months
    wb = __import__("openpyxl").Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Epyllion - 2026")
    ws.append(["Sl", "Suppliers", "Supplier", "Factory ", None, None, None])
    ws.cell(row=2, column=5, value=__import__("datetime").datetime(2026, 1, 1))
    ws.cell(row=3, column=1, value=1)
    ws.cell(row=3, column=2, value="SNQS GLOBAL TEXTILES FZE")
    ws.cell(row=3, column=3, value="SNQS GLOBAL TEXTILES FZE")
    ws.cell(row=3, column=4, value="Alim Knit Ltd (20331)")
    ws.cell(row=3, column=5, value=3949)
    path = tmp_path / "test.xlsx"
    wb.save(path)
    records, _ = load_records(path)
    assert len(records) == 1
    assert records[0].factory == "Alim Knit Ltd (20331)"
    assert records[0].supplier == "SNQS GLOBAL TEXTILES FZE"
    assert records[0].cartons == 3949
