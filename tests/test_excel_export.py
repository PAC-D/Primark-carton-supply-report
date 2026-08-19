import io

import openpyxl
import pandas as pd

from excel_export import render_excel
from filtering import month_label


def make_df(rows=5, months=3):
    periods = [(2024 + (8 + i - 1) // 12, (8 + i - 1) % 12 + 1) for i in range(months)]
    cols = ["SL", "Packaging Supplier", "Supplier", "Factory"] + [month_label(p) for p in periods] + ["Total"]
    data = []
    for i in range(rows):
        vals = [i + 1] * months
        data.append([i + 1, "M&U", f"Supplier {i % 10}", f"Factory {i}", *vals, sum(vals)])
    return pd.DataFrame(data, columns=cols)


def test_excel_export_creates_workbook(tmp_path):
    path = tmp_path / "out.xlsx"
    render_excel(make_df(), "M&U — Jan-26 to Mar-26", path)
    assert path.exists()
    assert path.stat().st_size > 1000


def test_excel_export_structure(tmp_path):
    path = tmp_path / "out.xlsx"
    df = make_df(rows=4, months=3)
    render_excel(df, "M&U — Jan-26 to Mar-26", path)
    ws = openpyxl.load_workbook(path)["Report"]
    assert ws["A1"].value == "M&U — Jan-26 to Mar-26"
    assert any(str(r).startswith("A1:") for r in ws.merged_cells.ranges)
    headers = [ws.cell(row=2, column=c).value for c in range(1, 9)]
    assert headers == df.columns.tolist()
    assert ws.freeze_panes == "A3"
    total_row = [ws.cell(row=7, column=c).value for c in range(1, 9)]
    assert total_row == [None, "Total", None, None, 10, 10, 10, 30]


def test_excel_export_styling(tmp_path):
    path = tmp_path / "out.xlsx"
    df = make_df(rows=4, months=3)
    render_excel(df, "M&U — Jan-26 to Mar-26", path)
    ws = openpyxl.load_workbook(path)["Report"]
    assert ws["A1"].font.bold
    assert ws["A1"].font.color.rgb.endswith("00205B")
    assert ws["A2"].font.bold
    assert ws["A2"].font.color.rgb.endswith("FFFFFF")
    assert ws["A2"].fill.start_color.rgb.endswith("00205B")
    assert ws["B7"].value == "Total"
    assert ws["B7"].font.bold
    assert ws["B7"].font.color.rgb.endswith("00205B")
    assert ws["B7"].fill.start_color.rgb.endswith("D9E2F3")
    assert ws["A3"].border.left.style == "thin"
    values = df["Supplier"].astype(str).tolist() + ["Supplier", ""]
    expected = min(max(len(v) for v in values) + 2, 40)
    assert ws.column_dimensions["C"].width == expected


def test_excel_export_whole_numbers(tmp_path):
    path = tmp_path / "out.xlsx"
    df = make_df(rows=3, months=2)
    df[df.columns[4]] = [1.0, 2.5, 3.9]
    df[df.columns[-1]] = [1.0, 2.5, 3.9]
    render_excel(df, "T", path)
    ws = openpyxl.load_workbook(path)["Report"]
    jan = [ws.cell(row=r, column=5).value for r in range(3, 6)]
    assert jan == [1, 2, 4]
    assert all(isinstance(v, int) for v in jan)


def test_excel_export_accepts_file_object():
    buf = io.BytesIO()
    render_excel(make_df(rows=3, months=2), "T", buf)
    assert len(buf.getvalue()) > 1000